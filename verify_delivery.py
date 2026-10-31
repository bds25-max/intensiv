"""Сверка поставки с исходными Excel. Запуск: python verify_delivery.py.

Создаёт карту фактической заполненности, сохраняет исходные значения истории
и исключённых колонок, сверяет нормализованную историю и повторяемость экспорта.
Исходные Excel не изменяются.
"""
import hashlib
import json
import tempfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import parser_core as core


OUTPUT_FILES = (
    "normalized_clients.json", "normalized_slices.jsonl", "rejected_slices.jsonl",
    "normalized_history.jsonl", "quality_report.json", "quality_report.csv",
)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def encode(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def exact_count(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).strip().replace(" ", "").replace(",", "."))
        return int(number) if number.is_finite() and number == number.to_integral_value() else None
    except (InvalidOperation, ValueError):
        return None


def load_jsonl(path):
    with Path(path).open(encoding="utf-8") as stream:
        return [json.loads(line) for line in stream if line.strip()]


def run():
    root = Path(__file__).resolve().parent
    sources = root / "sources"
    before = {p.name: sha256(p) for p in sorted(sources.iterdir()) if p.is_file()}
    clients = json.loads((root / OUTPUT_FILES[0]).read_text(encoding="utf-8"))
    valid = load_jsonl(root / OUTPUT_FILES[1])
    rejected = load_jsonl(root / OUTPUT_FILES[2])
    history = load_jsonl(root / OUTPUT_FILES[3])
    issues = json.loads((root / OUTPUT_FILES[4]).read_text(encoding="utf-8"))
    ids = {c["source_client_id"] for c in clients}
    check(len(ids) == len(clients), "Дубли карточек клиента")
    for row in clients + valid + history:
        check(row["source_client_id"] in ids, "Неизвестный клиент в основном результате")
    check(all(row["rejection_reason"] for row in rejected), "Отклонение без причины")
    keys = [(s["source_client_id"], s["period_start"], s["period_end"], s["as_of_date"]) for s in valid]
    check(len(keys) == len(set(keys)), "Повтор ключа среди допустимых срезов")

    # Независимый индекс истории по координатам, а не только контроль количества.
    history_index = {}
    for record in history:
        p = record["provenance"]
        key = (p["source_file"], p["source_sheet"], p["source_row"], p["source_column"], record["event_type"])
        check(key not in history_index, "Повтор координат истории")
        history_index[key] = record
    expected_history = set()
    checked_fields = 0
    coverage = []
    preserved = []
    formula_cells = 0
    uncached_formula_cells = 0
    books = {}
    source_sheet_keys = set()
    for filename in sorted({key.split("::", 1)[0] for key in core.COLUMN_MAPPING}):
        books[filename] = openpyxl.load_workbook(sources / filename, data_only=True)
        source_sheet_keys.update(filename + "::" + s for s in books[filename].sheetnames)
    check(source_sheet_keys == set(core.COLUMN_MAPPING), "Карта не совпадает с набором листов")

    def expect(file, sheet, row, col, event, fields):
        nonlocal checked_fields
        key = (file, sheet, row, col, event)
        check(key in history_index, "Нет записи истории: " + str(key))
        expected_history.add(key)
        for name, value in fields.items():
            check(history_index[key][name] == value, "Расхождение поля истории: " + str((key, name)))
            checked_fields += 1

    try:
        for sheet_key, columns in core.COLUMN_MAPPING.items():
            filename, sheetname = sheet_key.split("::", 1)
            ws = books[filename][sheetname]
            header_row = core.detect_header_row(ws)
            check([c["col_index"] for c in columns] == list(range(1, ws.max_column + 1)),
                  "Пропуск колонок карты: " + sheet_key)
            formulas = openpyxl.load_workbook(sources / filename, data_only=False)
            try:
                ws_f = formulas[sheetname]
                for column in columns:
                    c = column["col_index"]
                    actual_header = ws.cell(header_row, c).value
                    check(column["header"] == actual_header, "Заголовок карты отличается: " + sheet_key + column["col_letter"])
                    count = sum(text(ws.cell(r, c).value) is not None for r in range(header_row + 1, ws.max_row + 1))
                    coverage.append(dict(column, source_file=filename, source_sheet=sheetname,
                                         actual_header_row=header_row, actual_nonempty_cells=count))
                    target = column["target_field"]
                    for r in range(header_row + 1, ws.max_row + 1):
                        if ws_f.cell(r, c).data_type == "f":
                            formula_cells += 1
                            uncached_formula_cells += ws.cell(r, c).value is None
                    if target is not None and not target.startswith("HistoryRecord"):
                        continue
                    for r in range(header_row + 1, ws.max_row + 1):
                        original = ws_f.cell(r, c)
                        cached = ws.cell(r, c).value
                        if original.value is None and cached is None:
                            continue
                        preserved.append({
                            "source_client_id": text(ws.cell(r, 1).value),
                            "provenance": {"source_file": filename, "source_sheet": sheetname,
                                           "source_row": r, "source_column": column["col_letter"]},
                            "header": actual_header,
                            "category": "excluded_from_main_schema" if target is None else "history_source",
                            "raw_value": encode(original.value), "cached_value": encode(cached),
                            "excel_data_type": original.data_type,
                        })
            finally:
                formulas.close()

            if filename.startswith("КЦР_"):
                status_specs = [(6, "2026-07-21"), (7, "2026-05-21")]
                groups = []
            elif filename.startswith("База должники"):
                status_specs = [(c, dt) for c, _, dt in core.HistoricalFilesParser.STATUS_COLUMNS_SPEC]
                groups = [(46, "AT-AX", "first_contact"), (53, "BA-BE", "second_contact")]
            elif filename.startswith("База I кв"):
                status_specs = [(15, "2026-05-21")]
                groups = [(35, "AI-AM", "operator_contact")]
            else:
                status_specs, groups = [], []
            for r in range(header_row + 1, ws.max_row + 1):
                cid = text(ws.cell(r, 1).value)
                if cid not in ids:
                    continue
                for c, dt in status_specs:
                    value = text(ws.cell(r, c).value)
                    if value is not None:
                        expect(filename, sheetname, r, openpyxl.utils.get_column_letter(c), "historical_status",
                               {"source_client_id": cid, "occurred_at": dt, "raw_status": value})
                for start, col, event in groups:
                    values = [text(ws.cell(r, c).value) for c in range(start, start + 5)]
                    if any(v is not None for v in values):
                        expect(filename, sheetname, r, col, event, dict(
                            source_client_id=cid, occurred_at=None, text=values[0], raw_status=values[1],
                            raw_next_contact=values[2], raw_send_deadline=values[3], raw_client_date=values[4]))
                if filename.startswith("База должники"):
                    claim = text(ws.cell(r, 10).value)
                    if claim is not None:
                        expect(filename, sheetname, r, "J", "claim_comment",
                               {"source_client_id": cid, "text": claim, "occurred_at": None})
                    values = [text(ws.cell(r, c).value) for c in (58, 59, 60)]
                    if any(v is not None for v in values):
                        combined = "; ".join(label + v for label, v in zip(
                            ("Инфо от клиента: ", "Инцидент: ", "Запрос эксперта: "), values) if v is not None)
                        expect(filename, sheetname, r, "BF-BH", "incident_or_expert_request",
                               {"source_client_id": cid, "text": combined, "raw_deadline": values[0],
                                "raw_status": values[2], "occurred_at": None})
        check(expected_history == set(history_index), "Лишние или отсутствующие записи истории")

        # Координаты каждого основного объекта и происхождение атрибутов.
        for row in clients + valid + rejected + history:
            p = row["provenance"]
            ws = books[p["source_file"]][p["source_sheet"]]
            check(text(ws.cell(p["source_row"], 1).value) == row["source_client_id"],
                  "Идентификатор не совпадает с исходной строкой")
        for client in clients:
            for attr, p in client["attribute_provenance"].items():
                ws = books[p["source_file"]][p["source_sheet"]]
                raw = ws[p["source_column"] + str(p["source_row"])].value
                if attr == "creation_date":
                    expected = core.parse_iso_date(raw)
                elif attr == "raw_creation_date":
                    expected = raw if isinstance(raw, str) else text(raw)
                else:
                    expected = text(raw)
                check(client[attr] == expected, "Атрибут не совпадает с происхождением: " + attr)
        for record in valid + rejected:
            p = record["provenance"]
            ws = books[p["source_file"]][p["source_sheet"]]
            columns = p["source_column"].split(",")
            check(len(columns) == 2, "Нет координат количеств среза")
            for field, col in zip(("total", "unsent"), columns):
                value = ws[col + str(p["source_row"])].value
                check(record[field] == exact_count(value), "Количество не совпадает с исходной ячейкой")
    finally:
        for book in books.values():
            book.close()

    # Новый запуск и экспорт в отдельную временную папку, затем побайтная сверка.
    rerun = core.parse_all_kcr(str(sources))
    with tempfile.TemporaryDirectory() as temp:
        core.export_artifacts(*rerun, output_dir=temp)
        for filename in OUTPUT_FILES:
            check((Path(temp) / filename).read_bytes() == (root / filename).read_bytes(),
                  "Повторный запуск изменил " + filename)
    check(before == {p.name: sha256(p) for p in sorted(sources.iterdir()) if p.is_file()},
          "Исходные файлы изменились")
    result = {
        "status": "passed", "source_sha256": before, "sheets": len(source_sheet_keys),
        "columns": len(coverage), "clients": len(clients), "valid_slices": len(valid),
        "rejected_slices": len(rejected), "history_records": len(history), "quality_issues": len(issues),
        "history_fields_compared": checked_fields, "all_history_records_compared": True,
        "repeat_export_byte_identical": True, "sources_unchanged": True,
        "history_by_type": dict(Counter(r["event_type"] for r in history)),
        "issues_by_code": dict(Counter(r["code"] for r in issues)),
        "preserved_source_cells": len(preserved),
        "excluded_nonempty_cells": sum(r["category"] == "excluded_from_main_schema" for r in preserved),
        "formula_cells": formula_cells, "uncached_formula_cells": uncached_formula_cells,
        "slice_count_fields_compared": 2 * (len(valid) + len(rejected)),
    }
    for name, data in (("column_coverage.json", coverage), ("verification_result.json", result)):
        (root / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    with (root / "source_preserved_values.jsonl").open("w", encoding="utf-8") as stream:
        for record in preserved:
            stream.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + "\n")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    run()
